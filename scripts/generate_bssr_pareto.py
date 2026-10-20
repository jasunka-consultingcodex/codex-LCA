"""Build BSSR LCI and process Pareto workbooks in the established LMO layouts.

Usage: python scripts/generate_bssr_pareto.py SOURCE PROC_TEMPLATE LCI_TEMPLATE OUTPUT
       --lci-output LCI_OUTPUT --proc-output PROC_OUTPUT
"""

from __future__ import annotations

import argparse
import math
from pathlib import Path

from openpyxl import load_workbook

from generate_lmo_five_sheets import audit, parse_sheet, pareto, render_lci, render_proc
from generate_proc_dwn import ORDER


SHEETS = [
    (f"{kind} {phase} BSSR", kind, phase)
    for phase in ("HULU", "INTI", "UT")
    for kind in ("LCI", "PROC")
]


class SourceRows:
    def __init__(self, title, rows):
        self.title = title
        self._rows = rows

    @property
    def values(self):
        return iter(self._rows)


def recover_ut_water_inventory(source):
    """Read the LCI Water use block placed before PROC UT's process block."""
    rows = list(source["PROC UT BSSR"].values)
    starts = [i for i, row in enumerate(rows) if row[1] == "Category: " and row[2] == "Water use"]
    if len(starts) != 2:
        raise ValueError(f"Expected one inventory and one process Water use block, got {len(starts)}")
    headers = []
    for start, end in zip(starts, starts[1:] + [len(rows)]):
        header = next((row for row in rows[start:end] if row[1] == "No"), None)
        if header is None:
            raise ValueError("Water use block has no table header")
        headers.append(header[2])
    if headers != ["Substance", "Process"]:
        raise ValueError(f"Unexpected Water use table types: {headers}")
    # Preserve original row numbers in the audit, while parsing only the misplaced block.
    virtual_rows = [(None, None, None)] * starts[0] + rows[starts[0]:starts[1]]
    stages, groups = parse_sheet(SourceRows("PROC UT BSSR / inventory Water use", virtual_rows),
                                 "LCI", expected_impacts={"WSF"})
    for item in groups["WSF"]["items"].values():
        item["source_refs"] = [f"PROC UT BSSR!{row}" for row in sorted(item["rows"])]
    return stages, groups["WSF"]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("source", "proc_template", "lci_template", "output"):
        parser.add_argument(name, type=Path)
    parser.add_argument("--lci-output", type=Path, required=True)
    parser.add_argument("--proc-output", type=Path, required=True)
    args = parser.parse_args()
    outputs = [args.output, args.lci_output, args.proc_output]
    if len(set(outputs)) != len(outputs):
        raise ValueError("Output paths must be distinct")

    source = load_workbook(args.source, read_only=True, data_only=True)
    parsed, selections = {}, {}
    water_stages, water_group = recover_ut_water_inventory(source)
    for name, kind, phase in SHEETS:
        if name not in source:
            raise ValueError(f"Required source sheet missing: {name}")
        if name == "LCI UT BSSR":
            stages, groups = parse_sheet(source[name], kind, expected_impacts=set(ORDER) - {"WSF"})
            if stages != water_stages:
                raise ValueError("The misplaced UT water inventory has different unit-process columns")
            groups["WSF"] = water_group
        elif name == "PROC UT BSSR":
            stages, groups = parse_sheet(source[name], kind, skip_other_kind=True)
            if stages != water_stages:
                raise ValueError("UT water inventory and process unit columns differ")
        else:
            stages, groups = parse_sheet(source[name], kind)
        parsed[name] = (stages, groups)
        selections[name] = pareto(groups)
        print(f"{name}: {len(groups)} categories; {sum(map(len, selections[name].values()))} Pareto rows")
    for phase in ("HULU", "INTI", "UT"):
        proc_groups = parsed[f"PROC {phase} BSSR"][1]
        lci_groups = parsed[f"LCI {phase} BSSR"][1]
        for impact in ORDER:
            if (proc_groups[impact]["sources"] != lci_groups[impact]["sources"]
                    or not math.isclose(proc_groups[impact]["total"], lci_groups[impact]["total"],
                                        rel_tol=1e-6, abs_tol=1e-12)):
                raise ValueError(f"PROC/LCI indicators or totals differ: {phase} / {impact}")

    workbook = load_workbook(args.proc_template)
    proc_template = workbook["Hotspot LMO"]
    lci_template = load_workbook(args.lci_template)["LCIA - LCI LMO"]
    for name, kind, phase in SHEETS:
        stages, groups = parsed[name]
        if kind == "PROC":
            render_proc(workbook, proc_template, name, phase, stages, groups, selections[name], product="BSSR")
        else:
            render_lci(workbook, lci_template, name, phase, groups, selections[name],
                       parsed[f"PROC {phase} BSSR"][1], product="BSSR",
                       source_note=("Khusus WSF: blok inventori Water use berada di PROC UT BSSR baris 299–326; "
                                    "nomor baris rinci tercatat di Audit Pareto." if phase == "UT" else None))
    for sheet in list(workbook):
        if sheet.title not in {name for name, _, _ in SHEETS}:
            workbook.remove(sheet)
    workbook._sheets = [workbook[name] for name, _, _ in SHEETS]
    audit(workbook, parsed, selections, SHEETS, product="BSSR")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Saved {args.output}")

    for output, kind in ((args.lci_output, "LCI"), (args.proc_output, "PROC")):
        subset = load_workbook(args.output)
        names = [name for name, sheet_kind, _ in SHEETS if sheet_kind == kind]
        for sheet in list(subset):
            if sheet.title not in names:
                subset.remove(sheet)
        subset._sheets = [subset[name] for name in names]
        output.parent.mkdir(parents=True, exist_ok=True)
        subset.save(output)
        print(f"Saved {output}")


if __name__ == "__main__":
    main()
