"""Rebuild seven LMO Pareto sheets from a SimaPro export and both layout templates.

Usage: python scripts/generate_lmo_five_sheets.py SOURCE PROC_TEMPLATE LCI_TEMPLATE OUTPUT
"""

from __future__ import annotations

import argparse
import math
import re
from collections import defaultdict
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from generate_proc_dwn import ORDER, HEADERS, category, numeric, set_style


SHEETS = [
    ("LCI HULU LMO", "LCI", "HULU"),
    ("PROC HULU LMO", "PROC", "HULU"),
    ("LCI INTI LMO", "LCI", "INTI"),
    ("PROC INTI LMO", "PROC", "INTI"),
    ("LCI UT LMO", "LCI", "UT"),
    ("PROC UT LMO", "PROC", "UT"),
    ("PROC DOWN LMO", "PROC", "DOWN"),
]
LABELS = {
    "GWP": "Potensi pemanasan global", "ODP": "Potensi penipisan lapisan ozon",
    "AP": "Potensi hujan asam", "EP": "Potensi eutrofikasi",
    "POFP": "Photochemical oxidation", "ABF": "Abiotic depletion (fossil fuels)",
    "ABN": "Abiotic depletion", "ETP Terrestrial": "Ecotoxicity — Terrestrial",
    "ETP Freshwater": "Ecotoxicity — Freshwater", "ETP Marine": "Ecotoxicity — Marine",
    "HTC": "Human toxicity (karsinogenik)", "HTN": "Human toxicity (non-karsinogenik)",
    "WSF": "Water scarcity footprint", "LUC": "Land use change",
    "CEN": "Cumulative energy demand — non-renewable",
    "CED": "Cumulative energy demand — renewable",
}


def parse_sheet(sheet, kind):
    rows = list(sheet.values)
    starts = [(i, row[2]) for i, row in enumerate(rows) if row[1] == "Category: "]
    if not starts:
        raise ValueError(f"No categories in {sheet.title}")
    groups = defaultdict(lambda: {"total": 0.0, "unit": set(), "items": defaultdict(lambda: {"value": 0.0, "rows": set()}), "sources": []})
    seen, stages = set(), None
    for block, (start, raw) in enumerate(starts):
        if raw in seen:
            raise ValueError(f"Duplicate category in {sheet.title}: {raw}")
        seen.add(raw)
        end = starts[block + 1][0] if block + 1 < len(starts) else len(rows)
        header_at = next((i for i in range(start, end) if rows[i][1] == "No"), None)
        if header_at is None:
            raise ValueError(f"No table for {raw} in {sheet.title}")
        header = rows[header_at]
        right_no = next(i for i in range(2, len(header)) if header[i] == "No")
        shift = right_no - 1
        stage_columns = [(i, str(name)) for i, name in enumerate(header) if 6 <= i < right_no and isinstance(name, str) and re.match(r"^\d+\.", name)]
        stage_names = [name for _, name in stage_columns]
        if stages is None:
            stages = stage_names
        elif stages != stage_names:
            raise ValueError(f"Stage columns differ at {raw} in {sheet.title}")
        total_row = rows[header_at + 1]
        expected = "Total of all processes" if kind == "PROC" else "Total of all compartments"
        if total_row[2] != expected or not numeric(total_row[5 + shift]):
            raise ValueError(f"Invalid total for {raw} in {sheet.title}")
        total = total_row[5 + shift]
        if total == 0:
            raise ValueError(f"Zero total for {raw} in {sheet.title}")
        group = groups[category(raw)]
        group["total"] += total
        group["unit"].add(str(total_row[4 + shift]))
        group["sources"].append(raw)
        for row_index in range(header_at + 2, end):
            source_row = rows[row_index]
            if not numeric(source_row[1]) or not source_row[2]:
                continue
            for col, stage in stage_columns:
                amount = source_row[col + shift]
                if not numeric(amount) or abs(amount) < 1e-30:
                    continue
                percent = source_row[col]
                if numeric(percent) and abs(percent - 100 * amount / abs(total)) > 0.005:
                    raise ValueError(f"Left/right mismatch at {sheet.title}!{row_index + 1} / {raw}")
                if kind == "PROC":
                    key = (stage, str(source_row[2]))
                else:
                    key = (stage, str(source_row[2]), str(source_row[3] or ""))
                item = group["items"][key]
                item["value"] += amount
                item["rows"].add(row_index + 1)
    expected_impacts = set(ORDER)
    if set(groups) != expected_impacts:
        raise ValueError(f"Unexpected categories in {sheet.title}: {sorted(groups)}")
    for impact, group in groups.items():
        if len(group["unit"]) != 1:
            raise ValueError(f"Mixed units in {sheet.title} / {impact}: {group['unit']}")
    for impact in ("CEN", "CED"):
        if len(groups[impact]["sources"]) != 3:
            raise ValueError(f"Expected three energy subcategories for {impact} in {sheet.title}")
    return stages, groups


def pareto(groups):
    selected = {}
    for impact in ORDER:
        if impact not in groups:
            continue
        group = groups[impact]
        ranked = sorted(
            ((key, item, item["value"] / group["total"]) for key, item in group["items"].items()
             if item["value"] / group["total"] > 1e-12),
            key=lambda entry: (-entry[2], entry[0]),
        )
        rows, cumulative = [], 0.0
        for key, item, directional in ranked:
            cumulative += directional
            rows.append((key, item, item["value"] / abs(group["total"]), cumulative))
            if cumulative >= 0.8:
                break
        if not rows or rows[-1][3] < 0.8:
            raise ValueError(f"Pareto 80% not reached for {impact}")
        selected[impact] = rows
    return selected


def style_from(target, sample):
    target.font = copy(sample.font)
    target.fill = copy(sample.fill)
    target.border = copy(sample.border)
    target.alignment = copy(sample.alignment)
    target.protection = copy(sample.protection)
    target.number_format = sample.number_format


def render_proc(workbook, template, name, phase, stages, groups, selected):
    sheet = workbook.copy_worksheet(template)
    sheet.title = name
    for region in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(region))
    for row in sheet:
        for cell in row:
            cell.value = None
    if sheet.max_row > 7:
        sheet.delete_rows(8, sheet.max_row - 7)
    sheet.column_dimensions["B"].width = 35
    sheet.column_dimensions["C"].width = 65
    for col in range(4, 20):
        sheet.column_dimensions[get_column_letter(col)].width = 15
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D8"
    for row in (2, 3):
        for col in range(2, 20):
            set_style(sheet.cell(row, col), template["B2"])
    sheet.merge_cells("B2:S3")
    sheet["B2"] = f"HASIL ANALISIS ISU PENTING BERDASARKAN PROSES {phase} LMO\nKontribusi Pareto ≥80% per kategori dampak"
    sheet["B2"].alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    for row in (2, 3):
        sheet.row_dimensions[row].height = 25
    for col, value in ((2, "Unit Proses"), (3, "Proses")):
        for row in range(5, 8):
            set_style(sheet.cell(row, col), template.cell(5, col))
        sheet.cell(5, col, value)
        sheet.merge_cells(start_row=5, start_column=col, end_row=7, end_column=col)
    for col, impact in enumerate(ORDER, 4):
        set_style(sheet.cell(5, col), template["D5"])
        sheet.cell(5, col, HEADERS[col - 4])
        for row in (6, 7):
            set_style(sheet.cell(row, col), template["D6"])
        sheet.cell(6, col, next(iter(groups[impact]["unit"])) if impact in groups else "—")
        sheet.merge_cells(start_row=6, start_column=col, end_row=7, end_column=col)
        sheet.cell(5, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.cell(6, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row, height in ((5, 28), (6, 24), (7, 23)):
        sheet.row_dimensions[row].height = height
    for col in range(2, 20):
        set_style(sheet.cell(9, col), template["B32"])
    sheet.merge_cells("B9:S9")
    sheet["B9"] = f"Proses {phase} — kategori di kolom; nilai dampak di atas, kontribusi di bawah"
    sheet.row_dimensions[9].height = 23
    by_key = defaultdict(dict)
    for impact, entries in selected.items():
        for key, item, share, cumulative in entries:
            by_key[key][impact] = (item["value"], share)
    row = 10
    included = []
    for stage in stages:
        keys = [key for key in by_key if key[0] == stage]
        if not keys:
            continue
        included.append(stage)
        keys.sort(key=lambda key: (-max(abs(share) for _, share in by_key[key].values()), key[1]))
        first = row
        for key in keys:
            for col in range(2, 20):
                for current in (row, row + 1):
                    sample = template["B33" if col == 2 else "C33" if col == 3 else "D33"] if current == row else template["B34" if col == 2 else "C34" if col == 3 else "D34"]
                    set_style(sheet.cell(current, col), sample)
            sheet.cell(row, 3, key[1])
            sheet.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
            sheet.cell(row, 3).alignment = Alignment(vertical="center", wrap_text=True)
            for col, impact in enumerate(ORDER, 4):
                if impact in by_key[key]:
                    amount, share = by_key[key][impact]
                    sheet.cell(row, col, amount).number_format = "0.00E+00"
                    sheet.cell(row + 1, col, share).number_format = "0.00%"
            sheet.row_dimensions[row].height = 31
            sheet.row_dimensions[row + 1].height = 21
            row += 2
        sheet.cell(first, 2, stage)
        sheet.merge_cells(start_row=first, start_column=2, end_row=row - 1, end_column=2)
        sheet.cell(first, 2).alignment = Alignment(vertical="center", wrap_text=True)
    notes = [
        "Unit proses di luar Pareto: " + ("; ".join(stage for stage in stages if stage not in included) or "tidak ada"),
        f"Sumber: SimaPro LMO.xlsx / {'PROC DWN LMO' if phase == 'DOWN' else name}. Kontribusi bertanda mengikuti SimaPro; kumulatif arah neto ada di workbook gabungan / Audit Pareto. CEN/CED diagregasi; ETP terpisah.",
    ]
    for index, value in enumerate(notes, row + 1):
        sheet.merge_cells(start_row=index, start_column=2, end_row=index, end_column=19)
        cell = sheet.cell(index, 2, value)
        cell.font = Font(name="Aptos", size=10, italic=True, color="52616B")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
        sheet.row_dimensions[index].height = 27
    sheet.print_area = f"B2:S{row + 2}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    return sheet


def related_process(proc_groups, impact, stage):
    if impact not in proc_groups:
        return "Proses individual tidak tersedia pada sheet PROC"
    group = proc_groups[impact]
    candidates = [(key[1], item["value"] / group["total"]) for key, item in group["items"].items() if key[0] == stage and item["value"] / group["total"] > 0]
    if not candidates:
        return "Tidak ada proses searah hasil neto pada unit ini"
    return sorted(candidates, key=lambda pair: (-pair[1], pair[0]))[0][0]


def render_lci(workbook, template, name, phase, groups, selected, proc_groups):
    sheet = workbook.create_sheet(name)
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "F8"
    for col, width in {"B": 8, "C": 39, "D": 70, "E": 50, "F": 19, "G": 19, "H": 19, "I": 18}.items():
        sheet.column_dimensions[col].width = width
    for row in (2, 3):
        for col in range(2, 10):
            style_from(sheet.cell(row, col), template["B2"])
        sheet.row_dimensions[row].height = 25
    sheet.merge_cells("B2:I3")
    sheet["B2"] = f"HASIL ANALISIS ISU PENTING BERDASARKAN INVENTORY PRODUK — {phase} LMO\nKontribusi Pareto ≥80% per kategori dampak"
    sheet["B2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    headers = ["No", "Kategori Dampak", "Unit proses / proses terkait (bukan asal inventori)*", "Inventori / substansi penyebab", "Nilai dampak", "Satuan", "Kontribusi bertanda", "Kumulatif arah neto"]
    for col, heading in enumerate(headers, 2):
        for row in (5, 6, 7):
            style_from(sheet.cell(row, col), template["B5" if col <= 3 else "D5" if row == 5 else "E6"])
        sheet.cell(5, col, heading)
        sheet.cell(5, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=5, start_column=col, end_row=7, end_column=col)
    sheet.row_dimensions[5].height = 30
    sheet.row_dimensions[6].height = 20
    sheet.row_dimensions[7].height = 20
    row = 8
    sections = [("Dampak Primer", ORDER[:4]), ("Dampak Sekunder", ORDER[4:14]), ("Dampak Pemakaian Energi", ORDER[14:])]
    for section, impacts in sections:
        for col in range(2, 10):
            style_from(sheet.cell(row, col), template["B8"])
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
        sheet.cell(row, 2, section)
        sheet.row_dimensions[row].height = 24
        row += 1
        for impact in impacts:
            if impact not in groups:
                continue
            entries = selected[impact]
            for rank, (key, item, share, cumulative) in enumerate(entries, 1):
                for col in range(2, 10):
                    sample = template["B9"] if col == 2 else template["C9"] if col == 3 else template["D9"] if col == 4 else template["E9"]
                    style_from(sheet.cell(row, col), sample)
                    sheet.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
                stage, substance, compartment = key
                comparable = (impact in proc_groups
                              and proc_groups[impact]["sources"] == groups[impact]["sources"]
                              and math.isclose(proc_groups[impact]["total"], groups[impact]["total"], rel_tol=1e-6, abs_tol=1e-12))
                process = (related_process(proc_groups, impact, stage) if comparable
                           else "Proses tidak dipadankan: dasar indikator PROC dan LCI berbeda")
                values = [rank, f"{impact} — {LABELS[impact]}", f"{stage}\n{process}", f"{substance}\n[{compartment}]" if compartment else substance, item["value"], next(iter(groups[impact]["unit"])), share, cumulative]
                for col, value in enumerate(values, 2):
                    sheet.cell(row, col, value)
                sheet.cell(row, 6).number_format = "0.00E+00"
                for col in (8, 9):
                    sheet.cell(row, col).number_format = "0.00%"
                sheet.row_dimensions[row].height = 48 if len(process) > 90 or len(substance) > 55 else 36
                row += 1
    sheet.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=9)
    sheet.cell(row + 1, 2, "* Proses terkait ialah proses dengan kontribusi terbesar pada kategori dan unit yang sama di sheet PROC; data LCI tidak mengaitkan inventori langsung ke proses individual.")
    sheet.cell(row + 1, 2).alignment = Alignment(wrap_text=True, vertical="center")
    sheet.cell(row + 1, 2).font = Font(name="Aptos", size=10, italic=True, color="52616B")
    sheet.row_dimensions[row + 1].height = 35
    sheet.merge_cells(start_row=row + 2, start_column=2, end_row=row + 2, end_column=9)
    sheet.cell(row + 2, 2, f"Sumber: SimaPro LMO.xlsx / {name}. Nilai dihitung dari tabel kanan; kontribusi bertanda = nilai / |total|. CEN/CED diagregasi; ETP terpisah.")
    sheet.cell(row + 2, 2).font = Font(name="Aptos", size=10, italic=True, color="52616B")
    sheet.row_dimensions[row + 2].height = 25
    sheet.print_area = f"B2:I{row + 2}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3


def audit(workbook, parsed, selections):
    sheet = workbook.create_sheet("Audit Pareto")
    for col, width in enumerate([22, 23, 35, 65, 45, 20, 18, 18, 19, 21], 1):
        sheet.column_dimensions[get_column_letter(col)].width = width
    sheet.merge_cells("A1:J2")
    sheet["A1"] = "AUDIT PARETO — LMO / TUJUH SHEET"
    sheet["A1"].fill = PatternFill("solid", fgColor="244062")
    sheet["A1"].font = Font(name="Aptos", size=14, bold=True, color="FFFFFF")
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.merge_cells("A3:J3")
    sheet["A3"] = "Kumulatif = jumlah nilai/total neto; kontribusi bertanda = nilai/|total|. Baris berhenti pada pertama ≥80%."
    headers = ["Sheet sumber", "Kategori", "Unit proses", "Proses / substansi", "Kompartemen", "Nilai dampak", "Satuan", "Kontribusi", "Kumulatif", "Baris sumber"]
    for col, heading in enumerate(headers, 1):
        cell = sheet.cell(5, col, heading)
        cell.fill = PatternFill("solid", fgColor="DCEAF0")
        cell.font = Font(name="Aptos", size=10, bold=True, color="244062")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    row = 6
    for name, _, _ in SHEETS:
        _, groups = parsed[name]
        for impact in ORDER:
            if impact not in groups:
                continue
            group = groups[impact]
            for rank, (key, item, share, cumulative) in enumerate(selections[name][impact], 1):
                values = [name, impact, key[0], key[1], key[2] if len(key) > 2 else "", item["value"], next(iter(group["unit"])), share, cumulative, ", ".join(map(str, sorted(item["rows"])))]
                for col, value in enumerate(values, 1):
                    cell = sheet.cell(row, col, value)
                    cell.fill = PatternFill("solid", fgColor="E6F2E8" if rank == len(selections[name][impact]) else "F3F7F9" if row % 2 == 0 else "FFFFFF")
                    cell.font = Font(name="Aptos", size=10, color="202B32")
                    cell.alignment = Alignment(vertical="center", wrap_text=col in (3, 4, 5, 10))
                sheet.cell(row, 6).number_format = "0.00E+00"
                for col in (8, 9):
                    sheet.cell(row, col).number_format = "0.00%"
                sheet.row_dimensions[row].height = 30
                row += 1
    sheet.freeze_panes = "F6"
    sheet.sheet_view.showGridLines = False
    sheet.auto_filter.ref = f"A5:J{row - 1}"
    sheet.print_area = f"A1:J{row - 1}"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for arg in ("source", "proc_template", "lci_template", "output"):
        parser.add_argument(arg, type=Path)
    parser.add_argument("--lci-output", type=Path, help="Optional workbook containing only the LCI recap sheets")
    parser.add_argument("--proc-output", type=Path, help="Optional workbook containing only the PROC recap sheets")
    args = parser.parse_args()
    source = load_workbook(args.source, read_only=True, data_only=True)
    parsed = {}
    selections = {}
    for name, kind, _ in SHEETS:
        source_name = "PROC DWN LMO" if name == "PROC DOWN LMO" else name
        stages, groups = parse_sheet(source[source_name], kind)
        parsed[name] = (stages, groups)
        selections[name] = pareto(groups)
        print(f"{name}: {len(groups)} categories; {sum(map(len, selections[name].values()))} Pareto rows")
    workbook = load_workbook(args.proc_template)
    proc_template = workbook["Hotspot LMO"]
    lci_workbook = load_workbook(args.lci_template)
    lci_template = lci_workbook["LCIA - LCI LMO"]
    for name, kind, phase in SHEETS:
        stages, groups = parsed[name]
        if kind == "PROC":
            render_proc(workbook, proc_template, name, phase, stages, groups, selections[name])
        else:
            proc_name = f"PROC {phase} LMO"
            render_lci(workbook, lci_template, name, phase, groups, selections[name], parsed[proc_name][1])
    for sheet in list(workbook):
        if sheet.title not in [name for name, _, _ in SHEETS]:
            workbook.remove(sheet)
    workbook._sheets = [workbook[name] for name, _, _ in SHEETS]
    audit(workbook, parsed, selections)
    output_paths = [args.output] + [path for path in (args.lci_output, args.proc_output) if path is not None]
    if len(set(output_paths)) != len(output_paths):
        raise ValueError("Output paths must be distinct")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Saved {args.output}")
    for path, wanted in (
        (args.lci_output, {"LCI HULU LMO", "LCI INTI LMO", "LCI UT LMO"}),
        (args.proc_output, {"PROC HULU LMO", "PROC INTI LMO", "PROC UT LMO", "PROC DOWN LMO"}),
    ):
        if path is None:
            continue
        subset = load_workbook(args.output)
        for sheet in list(subset):
            if sheet.title not in wanted:
                subset.remove(sheet)
        path.parent.mkdir(parents=True, exist_ok=True)
        subset.save(path)
        print(f"Saved {path}")


if __name__ == "__main__":
    main()
