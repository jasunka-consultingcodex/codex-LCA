"""Build the LMO downstream Pareto workbook from SimaPro and the hotspot template.

Usage: python scripts/generate_proc_dwn.py SIMAPRO_XLSX TEMPLATE_XLSX OUTPUT_XLSX
"""

from __future__ import annotations

import argparse
import math
import re
from collections import defaultdict
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


ORDER = [
    "GWP", "ODP", "AP", "EP", "POFP", "ABF", "ABN",
    "ETP Terrestrial", "ETP Freshwater", "ETP Marine",
    "HTC", "HTN", "WSF", "LUC", "CEN", "CED",
]
HEADERS = [
    "GWP", "ODP", "AP", "EP", "POFP", "ABF", "ABN",
    "ETP Terra", "ETP Fresh", "ETP Marine", "HTC", "HTN",
    "WSF", "LUC", "CEN", "CED",
]


def category(raw: str) -> str:
    value = raw.lower()
    if value.startswith("gwp100"):
        return "GWP"
    if "ozone layer" in value:
        return "ODP"
    if value == "acidification":
        return "AP"
    if value == "eutrophication":
        return "EP"
    if value == "photochemical oxidation":
        return "POFP"
    if value == "abiotic depletion (fossil fuels)":
        return "ABF"
    if value == "abiotic depletion":
        return "ABN"
    if value == "terrestrial ecotoxicity":
        return "ETP Terrestrial"
    if value == "freshwater ecotoxicity":
        return "ETP Freshwater"
    if value == "marine ecotoxicity":
        return "ETP Marine"
    if value == "human carcinogenic toxicity":
        return "HTC"
    if value == "human non-carcinogenic toxicity":
        return "HTN"
    if value == "water use":
        return "WSF"
    if value == "land use":
        return "LUC"
    if value in ("non renewable, fossil", "non-renewable, nuclear", "non-renewable, biomass"):
        return "CEN"
    if value in ("renewable, biomass", "renewable, wind, solar, geothe", "renewable, water"):
        return "CED"
    raise ValueError(f"Unexpected impact category: {raw}")


def numeric(value: object) -> bool:
    return isinstance(value, (int, float)) and math.isfinite(value)


def extract(source: Path):
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet_name = "PROC DWN" if "PROC DWN" in workbook else "PROC DOWN"
    sheet = workbook[sheet_name]
    rows = list(sheet.values)
    starts = [(index, row[2]) for index, row in enumerate(rows) if row[1] == "Category: "]
    if not starts:
        raise ValueError(f"No SimaPro category blocks in {sheet_name}")

    groups = defaultdict(lambda: {"total": 0.0, "unit": set(), "items": defaultdict(lambda: {"value": 0.0, "rows": set()}), "sources": []})
    stages = []
    seen = set()
    for block, (start, raw_category) in enumerate(starts):
        end = starts[block + 1][0] if block + 1 < len(starts) else len(rows)
        if raw_category in seen:
            raise ValueError(f"Repeated source category needs inspection: {raw_category}")
        seen.add(raw_category)
        impact = category(raw_category)
        header_at = next((index for index in range(start, end) if rows[index][1] == "No"), None)
        if header_at is None:
            raise ValueError(f"Missing table for {raw_category}")
        header = rows[header_at]
        right_no = next(index for index in range(2, len(header)) if header[index] == "No")
        shift = right_no - 1
        stage_columns = [
            (column, name) for column, name in enumerate(header)
            if 6 <= column < right_no and isinstance(name, str) and re.match(r"^\d+\.", name)
        ]
        if not stages:
            stages = [name for _, name in stage_columns]
        elif stages != [name for _, name in stage_columns]:
            raise ValueError(f"Unit process columns changed at {raw_category}")

        total_row = rows[header_at + 1]
        if total_row[2] != "Total of all processes" or not numeric(total_row[5 + shift]):
            raise ValueError(f"Invalid total at {raw_category}")
        group = groups[impact]
        group["total"] += total_row[5 + shift]
        group["unit"].add(str(total_row[4 + shift]))
        group["sources"].append(raw_category)
        for row_index in range(header_at + 2, end):
            source_row = rows[row_index]
            if not numeric(source_row[1]) or not source_row[2]:
                continue
            for column, stage in stage_columns:
                amount = source_row[column + shift]
                if not numeric(amount) or abs(amount) < 1e-30:
                    continue
                source_percent = source_row[column]
                if numeric(source_percent) and abs(source_percent - 100 * amount / abs(total_row[5 + shift])) > 0.005:
                    raise ValueError(f"Percentage/value mismatch at {sheet_name}!{row_index + 1} ({raw_category})")
                item = group["items"][(stage, str(source_row[2]))]
                item["value"] += amount
                item["rows"].add(row_index + 1)

    if set(groups) != set(ORDER):
        raise ValueError(f"Expected 16 result categories; got {sorted(groups)}")
    for impact, group in groups.items():
        if group["total"] == 0 or len(group["unit"]) != 1:
            raise ValueError(f"Invalid total or units for {impact}")
    return sheet_name, stages, groups


def select_pareto(groups):
    result = {}
    for impact in ORDER:
        group = groups[impact]
        ranked = sorted(
            ((key, item, item["value"] / group["total"]) for key, item in group["items"].items()
             if item["value"] / group["total"] > 1e-12),
            key=lambda entry: (-entry[2], entry[0]),
        )
        chosen = []
        cumulative = 0.0
        for key, item, share in ranked:
            cumulative += share
            # SimaPro displays percentages against the magnitude of the total.
            # A negative net total therefore has negative signed percentages,
            # while Pareto coverage remains positive in the net direction.
            display_share = item["value"] / abs(group["total"])
            chosen.append((key, item, display_share, cumulative))
            if cumulative >= 0.8 - 1e-10:
                break
        if not chosen or chosen[-1][3] < 0.8:
            raise ValueError(f"Insufficient source contributions for {impact}")
        result[impact] = chosen
    return result


def set_style(target, sample):
    target._style = copy(sample._style)
    target.alignment = copy(sample.alignment)


def make_hotspot_sheet(workbook, template, stages, groups, selected):
    sheet = workbook.copy_worksheet(template)
    sheet.title = "Hotspot PROC DOWN"
    for region in list(sheet.merged_cells.ranges):
        sheet.unmerge_cells(str(region))
    for row in sheet:
        for cell in row:
            cell.value = None
    if sheet.max_row > 7:
        sheet.delete_rows(8, sheet.max_row - 7)

    sheet.column_dimensions["B"].width = 35
    sheet.column_dimensions["C"].width = 65
    for column in range(4, 20):
        sheet.column_dimensions[get_column_letter(column)].width = 15
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "D8"

    for row in (2, 3):
        for column in range(2, 20):
            set_style(sheet.cell(row, column), template["B2"])
    sheet.merge_cells("B2:S3")
    sheet["B2"] = "HASIL ANALISIS ISU PENTING BERDASARKAN PROSES HILIR LMO\nProduk sumber: 1 p Downstream — kontribusi Pareto ≥80% per kategori dampak"
    sheet["B2"].alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    sheet.row_dimensions[2].height = 25
    sheet.row_dimensions[3].height = 25

    for column, value in ((2, "Unit Proses"), (3, "Proses")):
        for row in range(5, 8):
            set_style(sheet.cell(row, column), template.cell(5, column))
        sheet.cell(5, column, value)
        sheet.merge_cells(start_row=5, start_column=column, end_row=7, end_column=column)
    for column, (impact, label) in enumerate(zip(ORDER, HEADERS), 4):
        set_style(sheet.cell(5, column), template["D5"])
        sheet.cell(5, column, label)
        unit = next(iter(groups[impact]["unit"]))
        for row in (6, 7):
            set_style(sheet.cell(row, column), template["D6"])
        sheet.cell(6, column, unit)
        sheet.merge_cells(start_row=6, start_column=column, end_row=7, end_column=column)
        sheet.cell(5, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.cell(6, column).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[5].height = 28
    sheet.row_dimensions[6].height = 24
    sheet.row_dimensions[7].height = 23

    for column in range(2, 20):
        set_style(sheet.cell(9, column), template["B32"])
    sheet.merge_cells("B9:S9")
    sheet["B9"] = "Proses Hilir (Downstream) — kategori dampak pada kolom; nilai dampak di atas, kontribusi di bawah"
    sheet.row_dimensions[9].height = 23

    chosen_by_key = defaultdict(dict)
    for impact, entries in selected.items():
        for key, item, share, cumulative in entries:
            chosen_by_key[key][impact] = (item["value"], share)
    row = 10
    included_stages = []
    for stage in stages:
        keys = [key for key in chosen_by_key if key[0] == stage]
        if not keys:
            continue
        included_stages.append(stage)
        keys.sort(key=lambda key: (-max(abs(share) for _, share in chosen_by_key[key].values()), key[1]))
        first = row
        for key in keys:
            name = key[1]
            for column in range(2, 20):
                set_style(sheet.cell(row, column), template["B33"] if column == 2 else template["C33"] if column == 3 else template["D33"])
                set_style(sheet.cell(row + 1, column), template["B34"] if column == 2 else template["C34"] if column == 3 else template["D34"])
            sheet.cell(row, 3, name)
            sheet.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
            sheet.cell(row, 3).alignment = Alignment(vertical="center", wrap_text=True)
            for column, impact in enumerate(ORDER, 4):
                if impact not in chosen_by_key[key]:
                    continue
                amount, share = chosen_by_key[key][impact]
                sheet.cell(row, column, amount).number_format = "0.00E+00"
                sheet.cell(row + 1, column, share).number_format = "0.00%"
            sheet.row_dimensions[row].height = 31
            sheet.row_dimensions[row + 1].height = 21
            row += 2
        sheet.cell(first, 2, stage)
        sheet.merge_cells(start_row=first, start_column=2, end_row=row - 1, end_column=2)
        sheet.cell(first, 2).alignment = Alignment(vertical="center", wrap_text=True)

    omitted = [stage for stage in stages if stage not in included_stages]
    sheet.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=19)
    note = sheet.cell(row + 1, 2)
    note.value = "Tidak masuk pilihan Pareto 80%: " + "; ".join(omitted) if omitted else "Semua unit proses muncul dalam pilihan Pareto."
    note.font = Font(name="Aptos", size=10, italic=True, color="52616B")
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row + 1].height = 27
    sheet.merge_cells(start_row=row + 2, start_column=2, end_row=row + 2, end_column=19)
    note = sheet.cell(row + 2, 2)
    note.value = "Sumber: SimaPro LMO.xlsx / PROC DWN. Persentase bertanda mengikuti SimaPro; kumulatif arah neto ada di Audit Pareto. CEN/CED masing-masing tiga subkategori; ETP tetap terpisah."
    note.font = Font(name="Aptos", size=10, italic=True, color="52616B")
    note.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[row + 2].height = 27
    sheet.print_area = f"B2:S{row + 2}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    return sheet


def make_audit_sheet(workbook, source_sheet, groups, selected):
    sheet = workbook.create_sheet("Audit Pareto")
    widths = [7, 25, 37, 82, 20, 19, 18, 18, 25]
    for column, width in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(column)].width = width
    navy, pale, light, green = "244062", "DCEAF0", "F3F7F9", "E6F2E8"
    sheet.merge_cells("A1:I2")
    title = sheet["A1"]
    title.value = "AUDIT PARETO — PROC DOWN / LMO"
    title.fill = PatternFill("solid", fgColor=navy)
    title.font = Font(name="Aptos", size=14, bold=True, color="FFFFFF")
    title.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 25
    sheet.row_dimensions[2].height = 20
    sheet.merge_cells("A3:I3")
    sheet["A3"] = f"Sumber: SimaPro LMO.xlsx / {source_sheet}. Kontribusi bertanda = nilai / |total|; berhenti pada kumulatif arah neto ≥80%."
    sheet.merge_cells("A4:I4")
    sheet["A4"] = "Kontribusi bertanda mengikuti SimaPro (nilai / |total|); kumulatif menghitung kontribusi searah hasil neto. Nilai >100% dapat terjadi karena aliran pengimbang."
    for row in (3, 4):
        sheet.cell(row, 1).font = Font(name="Aptos", size=10, italic=True, color="52616B")
    headers = ["No", "Kategori", "Unit Proses", "Proses", "Nilai dampak", "Satuan", "Kontribusi bertanda", "Kumulatif arah neto", "Baris sumber"]
    for column, name in enumerate(headers, 1):
        cell = sheet.cell(6, column, name)
        cell.fill = PatternFill("solid", fgColor=pale)
        cell.font = Font(name="Aptos", size=10, bold=True, color=navy)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[6].height = 27
    row = 7
    for impact in ORDER:
        group = groups[impact]
        for column in range(1, 10):
            cell = sheet.cell(row, column)
            cell.fill = PatternFill("solid", fgColor=pale)
            cell.font = Font(name="Aptos", size=10, bold=True, color=navy)
        sheet.cell(row, 2, impact)
        sheet.cell(row, 5, group["total"]).number_format = "0.00E+00"
        sheet.cell(row, 6, next(iter(group["unit"])))
        sheet.row_dimensions[row].height = 21
        row += 1
        for rank, (key, item, share, cumulative) in enumerate(selected[impact], 1):
            values = [rank, impact, key[0], key[1], item["value"], next(iter(group["unit"])), share, cumulative, ", ".join(map(str, sorted(item["rows"])))]
            final = rank == len(selected[impact])
            for column, value in enumerate(values, 1):
                cell = sheet.cell(row, column, value)
                cell.fill = PatternFill("solid", fgColor=green if final else light if row % 2 == 0 else "FFFFFF")
                cell.font = Font(name="Aptos", size=10, color="202B32")
                cell.alignment = Alignment(vertical="top", wrap_text=column in (3, 4, 9))
            sheet.cell(row, 5).number_format = "0.00E+00"
            for column in (7, 8):
                sheet.cell(row, column).number_format = "0.00%"
            sheet.row_dimensions[row].height = 35
            row += 1
        row += 1
    sheet.freeze_panes = "E7"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_view.zoomScale = 85
    sheet.print_area = f"A1:I{row - 1}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("template", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    source_sheet, stages, groups = extract(args.source)
    selected = select_pareto(groups)
    workbook = load_workbook(args.template)
    if "Hotspot LMO" not in workbook:
        raise ValueError("Template lacks Hotspot LMO")
    template = workbook["Hotspot LMO"]
    result = make_hotspot_sheet(workbook, template, stages, groups, selected)
    for sheet in list(workbook):
        if sheet is not result:
            workbook.remove(sheet)
    make_audit_sheet(workbook, source_sheet, groups, selected)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Saved {args.output} — {len(groups)} categories, {sum(map(len, selected.values()))} Pareto rows")


if __name__ == "__main__":
    main()
