"""Build GMO LCI and process Pareto workbooks in the established LMO layouts.

Usage: python scripts/generate_gmo_pareto.py SOURCE PROC_TEMPLATE LCI_TEMPLATE OUTPUT
       --lci-output LCI_OUTPUT --proc-output PROC_OUTPUT
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook

from generate_lmo_five_sheets import audit, parse_sheet, pareto, render_lci, render_proc


SHEETS = [
    ("LCI HULU GMO", "LCI", "HULU"),
    ("PROC HULU GMO", "PROC", "HULU"),
    ("LCI INTI GMO", "LCI", "INTI"),
    ("PROC INTI GMO", "PROC", "INTI"),
    ("LCI UT GMO", "LCI", "UT"),
    ("PROC UT GMO", "PROC", "UT"),
    ("PROC DWN GMO", "PROC", "DWN"),
]


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for name in ("source", "proc_template", "lci_template", "output"):
        parser.add_argument(name, type=Path)
    parser.add_argument("--lci-output", type=Path, required=True)
    parser.add_argument("--proc-output", type=Path, required=True)
    args = parser.parse_args()
    outputs = [args.output, args.lci_output, args.proc_output]
    if len(set(outputs)) != len(outputs):
        raise ValueError("Output paths must be distinct")

    source = load_workbook(args.source, read_only=True, data_only=True)
    parsed, selections = {}, {}
    for name, kind, _ in SHEETS:
        if name not in source:
            raise ValueError(f"Required source sheet missing: {name}")
        stages, groups = parse_sheet(source[name], kind)
        parsed[name] = (stages, groups)
        selections[name] = pareto(groups)
        print(f"{name}: {len(groups)} categories; {sum(map(len, selections[name].values()))} Pareto rows")

    workbook = load_workbook(args.proc_template)
    proc_template = workbook["Hotspot LMO"]
    lci_template = load_workbook(args.lci_template)["LCIA - LCI LMO"]
    for name, kind, phase in SHEETS:
        stages, groups = parsed[name]
        if kind == "PROC":
            render_proc(workbook, proc_template, name, phase, stages, groups, selections[name], product="GMO")
        else:
            render_lci(workbook, lci_template, name, phase, groups, selections[name],
                       parsed[f"PROC {phase} GMO"][1], product="GMO")
    for sheet in list(workbook):
        if sheet.title not in {name for name, _, _ in SHEETS}:
            workbook.remove(sheet)
    workbook._sheets = [workbook[name] for name, _, _ in SHEETS]
    audit(workbook, parsed, selections, SHEETS, product="GMO")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Saved {args.output}")

    for output, kind in ((args.lci_output, "LCI"), (args.proc_output, "PROC")):
        subset = load_workbook(args.output)
        names = [name for name, sheet_kind, _ in SHEETS if sheet_kind == kind]
        for sheet in list(subset):
            if sheet.title not in names:
                subset.remove(sheet)
        subset._sheets = [subset[name] for name in names]
        output.parent.mkdir(parents=True, exist_ok=True)
        subset.save(output)
        print(f"Saved {output}")


if __name__ == "__main__":
    main()
