"""Build the two-sheet BMO Pareto workbook using the supplied LMO workbook as layout.

Usage: python scripts/generate_bmo_pareto.py SIMAPRO_BMO_XLSX LMO_LAYOUT_XLSX OUTPUT_XLSX
"""

from __future__ import annotations

import argparse
import re
from collections import defaultdict
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.pagebreak import Break

from generate_lmo_five_sheets import HEADERS, LABELS, ORDER, category, numeric, pareto, related_process, style_from


PROC = ["PROC HULU BMO", "PROC INTI BMO", "PROC UT BMO", "PROC DWN BMO"]
LCI = ["LCI HULU BMO", "LCI INTI BMO", "LCI UT BMO"]
DISPLAY = {"HULU": "HULU", "INTI": "INTI", "UT": "UTILITAS", "DWN": "HILIR (DOWNSTREAM)"}


def parse_source(sheet, kind):
    rows = list(sheet.values)
    starts = [(i, r[2]) for i, r in enumerate(rows) if r[1] == "Category: "]
    if not starts:
        raise ValueError(f"Tidak ada blok kategori di {sheet.title}")
    groups = defaultdict(lambda: {"total": 0.0, "unit": set(), "items": defaultdict(lambda: {"value": 0.0, "rows": set()}), "sources": []})
    seen, stages = set(), None
    for block, (start, raw) in enumerate(starts):
        if raw in seen:
            raise ValueError(f"Kategori berulang: {sheet.title} / {raw}")
        seen.add(raw)
        end = starts[block + 1][0] if block + 1 < len(starts) else len(rows)
        header_at = next((i for i in range(start, end) if rows[i][1] == "No"), None)
        if header_at is None:
            raise ValueError(f"Tabel tidak ditemukan: {sheet.title} / {raw}")
        header = rows[header_at]
        right_no = next(i for i in range(2, len(header)) if header[i] == "No")
        shift = right_no - 1
        stage_columns = [(i, str(v)) for i, v in enumerate(header) if 6 <= i < right_no and isinstance(v, str) and re.match(r"^\d+\.", v)]
        names = [name for _, name in stage_columns]
        if stages is None:
            stages = names
        elif stages != names:
            raise ValueError(f"Unit proses berubah: {sheet.title} / {raw}")
        total_row = rows[header_at + 1]
        expected = "Total of all processes" if kind == "PROC" else "Total of all compartments"
        expected_header = "Process" if kind == "PROC" else "Substance"
        if header[2] != expected_header or total_row[2] != expected or not numeric(total_row[5 + shift]):
            raise ValueError(f"Total tidak valid: {sheet.title} / {raw}")
        total = total_row[5 + shift]
        if total == 0:
            raise ValueError(f"Total nol: {sheet.title} / {raw}")
        group = groups[category(raw)]
        group["total"] += total
        group["unit"].add(str(total_row[4 + shift]))
        group["sources"].append(raw)
        for i in range(header_at + 2, end):
            r = rows[i]
            if not numeric(r[1]) or not r[2]:
                continue
            for col, stage in stage_columns:
                amount = r[col + shift]
                if not numeric(amount) or abs(amount) < 1e-30:
                    continue
                percentage = r[col]
                if numeric(percentage) and abs(percentage - 100 * amount / abs(total)) > 0.005:
                    raise ValueError(f"Persentase dan nilai berbeda: {sheet.title}!{i + 1} / {raw}")
                key = (stage, str(r[2])) if kind == "PROC" else (stage, str(r[2]), str(r[3] or ""))
                item = group["items"][key]
                item["value"] += amount
                item["rows"].add(i + 1)
    if not set(groups).issubset(ORDER):
        raise ValueError(f"Kategori tak dikenal di {sheet.title}")
    for impact, group in groups.items():
        if len(group["unit"]) != 1:
            raise ValueError(f"Satuan campur: {sheet.title} / {impact}")
    for energy in ("CEN", "CED"):
        if energy in groups and len(groups[energy]["sources"]) != 3:
            raise ValueError(f"{sheet.title}: {energy} harus memiliki tiga subkategori")
    if set(groups) != set(ORDER):
        raise ValueError(f"{sheet.title}: kategori tidak lengkap ({sorted(groups)})")
    return stages, groups


def proc_table(sheet, template, start, source_name, stages, groups, chosen):
    phase = source_name.split()[1]
    # Each table repeats the title, green two-tier impact header, grey section
    # band and two-row process entries from the supplied PROC layout.
    for row in (start, start + 1):
        for col in range(2, 20):
            style_from(sheet.cell(row, col), template["B2"])
        sheet.row_dimensions[row].height = 25
    sheet.merge_cells(start_row=start, start_column=2, end_row=start + 1, end_column=19)
    title = sheet.cell(start, 2, f"HASIL ANALISIS ISU PENTING BERDASARKAN PROSES {DISPLAY[phase]} BMO\nPareto ≥80% per kategori dampak — sumber: {source_name}")
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top = start + 3
    for col, label in ((2, "Unit Proses"), (3, "Nama Kegiatan / Proses")):
        for row in range(top, top + 3):
            style_from(sheet.cell(row, col), template.cell(5, col))
        sheet.cell(top, col, label)
        sheet.merge_cells(start_row=top, start_column=col, end_row=top + 2, end_column=col)
    for col, impact in enumerate(ORDER, 4):
        style_from(sheet.cell(top, col), template["D5"])
        sheet.cell(top, col, HEADERS[col - 4])
        sheet.cell(top, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in (top + 1, top + 2):
            style_from(sheet.cell(row, col), template["D6"])
        sheet.cell(top + 1, col, next(iter(groups[impact]["unit"])) if impact in groups else "—")
        sheet.cell(top + 1, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=top + 1, start_column=col, end_row=top + 2, end_column=col)
    for row, height in ((top, 28), (top + 1, 24), (top + 2, 23)):
        sheet.row_dimensions[row].height = height
    band = top + 4
    for col in range(2, 20):
        style_from(sheet.cell(band, col), template["B9"])
    sheet.merge_cells(start_row=band, start_column=2, end_row=band, end_column=19)
    sheet.cell(band, 2, f"Proses {DISPLAY[phase]} — nilai dampak di atas; kontribusi bertanda di bawah")
    sheet.row_dimensions[band].height = 24

    by_key = defaultdict(dict)
    for impact, entries in chosen.items():
        for key, item, share, _ in entries:
            by_key[key][impact] = (item["value"], share)
    row = band + 1
    used_stages = []
    for stage in stages:
        keys = [key for key in by_key if key[0] == stage]
        if not keys:
            continue
        used_stages.append(stage)
        keys.sort(key=lambda key: (-max(abs(share) for _, share in by_key[key].values()), key[1]))
        first = row
        for key in keys:
            for col in range(2, 20):
                for dest, sample_row in ((row, 10), (row + 1, 11)):
                    sample = template.cell(sample_row, 2 if col == 2 else 3 if col == 3 else 4)
                    style_from(sheet.cell(dest, col), sample)
            sheet.cell(row, 3, key[1])
            sheet.merge_cells(start_row=row, start_column=3, end_row=row + 1, end_column=3)
            sheet.cell(row, 3).alignment = Alignment(vertical="center", wrap_text=True)
            for col, impact in enumerate(ORDER, 4):
                if impact in by_key[key]:
                    value, share = by_key[key][impact]
                    sheet.cell(row, col, value).number_format = "0.00E+00"
                    sheet.cell(row + 1, col, share).number_format = "0.00%"
            sheet.row_dimensions[row].height = 32
            sheet.row_dimensions[row + 1].height = 21
            row += 2
        sheet.cell(first, 2, stage)
        sheet.merge_cells(start_row=first, start_column=2, end_row=row - 1, end_column=2)
        sheet.cell(first, 2).alignment = Alignment(vertical="center", wrap_text=True)
    omitted = [stage for stage in stages if stage not in used_stages]
    sheet.merge_cells(start_row=row + 1, start_column=2, end_row=row + 1, end_column=19)
    note = sheet.cell(row + 1, 2, "Unit proses di luar pilihan Pareto: " + ("; ".join(omitted) or "tidak ada"))
    note.font = Font(name="Aptos", size=10, italic=True, color="52616B")
    note.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[row + 1].height = 25
    return row + 1


def lci_table(sheet, template, start, source_name, groups, chosen, proc_groups):
    phase = source_name.split()[1]
    for row in (start, start + 1):
        for col in range(2, 10):
            style_from(sheet.cell(row, col), template["B2"])
        sheet.row_dimensions[row].height = 25
    sheet.merge_cells(start_row=start, start_column=2, end_row=start + 1, end_column=9)
    title = sheet.cell(start, 2, f"HASIL ANALISIS ISU PENTING BERDASARKAN INVENTORY PRODUK — {DISPLAY[phase]} BMO\nPareto ≥80% per kategori dampak — sumber: {source_name}")
    title.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    top = start + 3
    headings = ["No", "Kategori Dampak", "Unit proses / proses terkait*", "Inventori / substansi penyebab", "Nilai dampak", "Satuan", "Kontribusi bertanda", "Kumulatif arah neto"]
    for col, heading in enumerate(headings, 2):
        for row in range(top, top + 3):
            sample = template["B5" if col <= 3 else "D5" if row == top else "E6"]
            style_from(sheet.cell(row, col), sample)
        sheet.cell(top, col, heading)
        sheet.cell(top, col).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=top, start_column=col, end_row=top + 2, end_column=col)
    for row, height in ((top, 30), (top + 1, 20), (top + 2, 20)):
        sheet.row_dimensions[row].height = height
    row = top + 3
    sections = [("Dampak Primer", ORDER[:4]), ("Dampak Sekunder", ORDER[4:14]), ("Dampak Pemakaian Energi", ORDER[14:])]
    for title, impacts in sections:
        available = [impact for impact in impacts if impact in groups]
        if not available:
            continue
        for col in range(2, 10):
            style_from(sheet.cell(row, col), template["B8"])
        sheet.merge_cells(start_row=row, start_column=2, end_row=row, end_column=9)
        sheet.cell(row, 2, title)
        sheet.row_dimensions[row].height = 24
        row += 1
        for impact in available:
            for rank, (key, item, share, cumulative) in enumerate(chosen[impact], 1):
                for col in range(2, 10):
                    sample = template["B9"] if col == 2 else template["C9"] if col == 3 else template["D9"] if col == 4 else template["E9"]
                    style_from(sheet.cell(row, col), sample)
                    sheet.cell(row, col).alignment = Alignment(vertical="center", wrap_text=True)
                stage, substance, compartment = key
                process = related_process(proc_groups, impact, stage) if impact in proc_groups else f"Proses individual tidak tersedia pada PROC {phase} BMO"
                values = [rank, f"{impact} — {LABELS[impact]}", f"{stage}\n{process}", f"{substance}\n[{compartment}]" if compartment else substance, item["value"], next(iter(groups[impact]["unit"])), share, cumulative]
                for col, value in enumerate(values, 2):
                    sheet.cell(row, col, value)
                sheet.cell(row, 6).number_format = "0.00E+00"
                for col in (8, 9):
                    sheet.cell(row, col).number_format = "0.00%"
                sheet.row_dimensions[row].height = 50 if len(process) > 90 or len(substance) > 55 else 37
                row += 1
    notes = [
        "* Proses terkait dipilih dari kategori dan unit yang sama pada sheet PROC; sumber LCI tidak menunjukkan hubungan langsung inventori dengan proses individual.",
        "Semua 16 kategori dampak tersedia. CEN/CED masing-masing menjumlahkan tiga subkategori; ETP tetap terpisah.",
    ]
    for offset, message in enumerate(notes, 1):
        sheet.merge_cells(start_row=row + offset, start_column=2, end_row=row + offset, end_column=9)
        cell = sheet.cell(row + offset, 2, message)
        cell.font = Font(name="Aptos", size=10, italic=True, color="52616B")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.row_dimensions[row + offset].height = 30
    return row + 2


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for arg in ("source", "template", "output"):
        parser.add_argument(arg, type=Path)
    args = parser.parse_args()
    source = load_workbook(args.source, read_only=True, data_only=True)
    if not set(PROC + LCI).issubset(source.sheetnames):
        raise ValueError("Sheet BMO yang diminta tidak lengkap")
    parsed, selected = {}, {}
    for name in PROC + LCI:
        stages, groups = parse_source(source[name], name.split()[0])
        parsed[name] = (stages, groups)
    for name in PROC + LCI:
        selected[name] = pareto(parsed[name][1])
        print(name, len(parsed[name][1]), "kategori,", sum(map(len, selected[name].values())), "baris Pareto")
    reference = load_workbook(args.template)
    proc_template = reference["PROC INTI LMO"]
    lci_template = reference["LCI INTI LMO"]
    output = Workbook()
    proc = output.active
    proc.title = "Hotspot Proses BMO"
    lci = output.create_sheet("LCIA - LCI BMO Pareto")
    for sheet in (proc, lci):
        sheet.sheet_view.showGridLines = False
        sheet.sheet_view.zoomScale = 80
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    proc.freeze_panes = "D8"
    proc.column_dimensions["B"].width = 35
    proc.column_dimensions["C"].width = 65
    for col in range(4, 20):
        proc.column_dimensions[get_column_letter(col)].width = 15
    row = 2
    for name in PROC:
        stages, groups = parsed[name]
        if row > 2:
            proc.row_breaks.append(Break(id=row - 1))
        row = proc_table(proc, proc_template, row, name, stages, groups, selected[name]) + 4
    proc.print_area = f"B2:S{row - 4}"
    for col, width in {"B": 8, "C": 39, "D": 70, "E": 50, "F": 19, "G": 19, "H": 19, "I": 19}.items():
        lci.column_dimensions[col].width = width
    lci.freeze_panes = "F8"
    row = 2
    for name in LCI:
        phase = name.split()[1]
        if row > 2:
            lci.row_breaks.append(Break(id=row - 1))
        row = lci_table(lci, lci_template, row, name, parsed[name][1], selected[name], parsed[f"PROC {phase} BMO"][1]) + 4
    lci.print_area = f"B2:I{row - 4}"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    output.save(args.output)
    print("Tersimpan:", args.output)


if __name__ == "__main__":
    main()
