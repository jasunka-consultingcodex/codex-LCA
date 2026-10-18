"""Build the SMO LCI and process Pareto workbooks using the LMO layout templates.

Usage: python scripts/generate_smo_pareto.py SOURCE PROC_TEMPLATE LCI_TEMPLATE OUTPUT
       --lci-output LCI_OUTPUT --proc-output PROC_OUTPUT
"""

from __future__ import annotations

import argparse
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font

from generate_lmo_five_sheets import (
    audit, parse_sheet, pareto, render_lci, render_proc, style_from,
)


SHEETS = [
    (f"{kind} {phase} SMO", kind, phase)
    for phase in ("HULU", "INTI", "UT", "DWN")
    for kind in ("LCI", "PROC")
]


def unavailable_lci(workbook, template):
    """Show the missing downstream source explicitly without fabricating inventory rows."""
    sheet = workbook.create_sheet("LCI DWN SMO")
    sheet.sheet_view.showGridLines = False
    for col, width in {"B": 8, "C": 39, "D": 70, "E": 50, "F": 19, "G": 19, "H": 19, "I": 18}.items():
        sheet.column_dimensions[col].width = width
    for row in (2, 3):
        for col in range(2, 10):
            style_from(sheet.cell(row, col), template["B2"])
        sheet.row_dimensions[row].height = 25
    sheet.merge_cells("B2:I3")
    sheet["B2"] = "HASIL ANALISIS ISU PENTING BERDASARKAN INVENTORY PRODUK — DWN SMO\nSumber LCI hilir belum tersedia"
    sheet["B2"].alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    headers = ["No", "Kategori Dampak", "Unit proses / proses terkait", "Inventori / substansi penyebab", "Nilai dampak", "Satuan", "Kontribusi bertanda", "Kumulatif arah neto"]
    for col, heading in enumerate(headers, 2):
        for row in (5, 6, 7):
            style_from(sheet.cell(row, col), template["B5" if col <= 3 else "D5" if row == 5 else "E6"])
        sheet.cell(5, col, heading).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.merge_cells(start_row=5, start_column=col, end_row=7, end_column=col)
    sheet.row_dimensions[5].height = 30
    sheet.row_dimensions[6].height = 20
    sheet.row_dimensions[7].height = 20
    for col in range(2, 10):
        style_from(sheet.cell(9, col), template["B8"])
    sheet.merge_cells("B9:I9")
    sheet["B9"] = "LCI DWN SMO — data inventori tidak tersedia"
    sheet.row_dimensions[9].height = 25
    sheet.merge_cells("B10:I12")
    sheet["B10"] = (
        "Berkas SimaPro SMO.xlsx tidak memuat sheet LCI DWN SMO. "
        "Kontribusi inventori dan Pareto hilir belum dapat dihitung. "
        "Sheet PROC DWN SMO tersedia dan direkap dalam workbook proses."
    )
    sheet["B10"].alignment = Alignment(vertical="center", wrap_text=True)
    sheet["B10"].font = Font(name="Aptos", size=11, color="202B32")
    sheet.row_dimensions[10].height = 30
    sheet.row_dimensions[11].height = 25
    sheet.row_dimensions[12].height = 25
    sheet.print_area = "B2:I12"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    for arg in ("source", "proc_template", "lci_template", "output"):
        parser.add_argument(arg, type=Path)
    parser.add_argument("--lci-output", type=Path, required=True)
    parser.add_argument("--proc-output", type=Path, required=True)
    args = parser.parse_args()
    output_paths = [args.output, args.lci_output, args.proc_output]
    if len(set(output_paths)) != len(output_paths):
        raise ValueError("Output paths must be distinct")

    source = load_workbook(args.source, read_only=True, data_only=True)
    parsed, selections, available = {}, {}, []
    for name, kind, phase in SHEETS:
        if name not in source:
            if name == "LCI DWN SMO":
                print(f"{name}: source sheet missing; a clearly marked placeholder will be created")
                continue
            raise ValueError(f"Required source sheet missing: {name}")
        stages, groups = parse_sheet(source[name], kind)
        parsed[name] = (stages, groups)
        selections[name] = pareto(groups)
        available.append((name, kind, phase))
        print(f"{name}: {len(groups)} categories; {sum(map(len, selections[name].values()))} Pareto rows")

    workbook = load_workbook(args.proc_template)
    proc_template = workbook["Hotspot LMO"]
    lci_template = load_workbook(args.lci_template)["LCIA - LCI LMO"]
    for name, kind, phase in SHEETS:
        if name not in parsed:
            unavailable_lci(workbook, lci_template)
            continue
        stages, groups = parsed[name]
        if kind == "PROC":
            render_proc(workbook, proc_template, name, phase, stages, groups, selections[name], product="SMO")
        else:
            proc_groups = parsed[f"PROC {phase} SMO"][1]
            render_lci(workbook, lci_template, name, phase, groups, selections[name], proc_groups, product="SMO")
    for sheet in list(workbook):
        if sheet.title not in {name for name, _, _ in SHEETS}:
            workbook.remove(sheet)
    workbook._sheets = [workbook[name] for name, _, _ in SHEETS]
    audit(workbook, parsed, selections, available, product="SMO")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    print(f"Saved {args.output}")

    for output, kind in ((args.lci_output, "LCI"), (args.proc_output, "PROC")):
        subset = load_workbook(args.output)
        names = [name for name, sheet_kind, _ in SHEETS if sheet_kind == kind]
        for sheet in list(subset):
            if sheet.title not in names:
                subset.remove(sheet)
        subset._sheets = [subset[name] for name in names]
        output.parent.mkdir(parents=True, exist_ok=True)
        subset.save(output)
        print(f"Saved {output}")


if __name__ == "__main__":
    main()
